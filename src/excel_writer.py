"""
V6 模板 Excel 生成器 — 基于 config/成衣工艺单_设计师填表模板_v6.xlsx 填数
"""
import os, sys, json
from pathlib import Path
from openpyxl import load_workbook

TEMPLATE = "成衣工艺单_设计师填表模板_v6.xlsx"

def _find_template():
    for base in [Path(__file__).resolve().parent.parent / "config",
                 Path(__file__).resolve().parent.parent / "skill" / "config",
                 Path.home() / "Desktop"]:
        p = base / TEMPLATE
        if p.exists():
            return p
    raise FileNotFoundError(f"模板未找到: {TEMPLATE}")

def build_excel(data: dict, output_path: str):
    wb = load_workbook(_find_template())
    ws = wb.active

    # === 基础信息 ===
    ws["B6"] = data.get("style_number", "")
    ws["D6"] = data.get("需求类型", data.get("requirement_type", "ODM"))
    ws["B7"] = data.get("衣服类型", data.get("garment_type", "T"))
    ws["D7"] = data.get("款式", data.get("style_name", ""))
    ws["B8"] = data.get("客户名称", data.get("customer", ""))
    ws["B9"] = int(data.get("打版件数", data.get("sample_quantity", 1)))
    ws["D9"] = data.get("版型", data.get("layout", ""))
    ws["B10"] = data.get("烫唛", data.get("mark_name", ""))
    ws["D10"] = data.get("预计完成时间", data.get("finish_time", ""))

    # === 工艺制作 (rows 14-17) ===
    process = data.get("工艺制作") or data.get("process_notes", "")
    if isinstance(process, str):
        lines = [l.strip() for l in process.replace("；", ";").split("\n") if l.strip()]
    elif isinstance(process, list):
        lines = process
    else:
        lines = []
    for i, line in enumerate(lines):
        r = 14 + i
        if r <= 17:
            ws.cell(r, 1, i + 1)
            ws.cell(r, 2, str(line))

    # === SKU + 面料 (rows 21-29) ===
    fabrics = data.get("fabrics") or data.get("面料信息", [])
    color = data.get("颜色") or data.get("sample_color", "")
    size = data.get("尺码") or data.get("sample_size", "L")
    qty = data.get("件数") or data.get("sample_quantity", 1)
    designer = data.get("designer_name") or data.get("designer", "")

    for idx, f in enumerate(fabrics):
        r = 21 + idx
        if r > 29:
            break
        if idx == 0:
            ws.cell(r, 1, 1)
            ws.cell(r, 2, color)
            ws.cell(r, 3, size)
            ws.cell(r, 4, int(qty))
        ws.cell(r, 5, f.get("code") or f.get("布类", chr(65 + idx)))
        ws.cell(r, 6, f.get("name") or f.get("面料编号", ""))
        ws.cell(r, 7, f.get("色号") or f.get("color", ""))
        ws.cell(r, 8, f.get("颜色名") or f.get("color_name", ""))
        ws.cell(r, 9, f.get("是否罗纹", ""))
        ws.cell(r, 10, f.get("width") or f.get("幅宽", ""))
        ws.cell(r, 11, f.get("weight") or f.get("克重", ""))
        ws.cell(r, 12, f.get("supplier") or f.get("供应商", "嘉谦纺织"))
        ws.cell(r, 13, f.get("component") or f.get("成分", ""))
        ws.cell(r, 14, f.get("product_name") or f.get("品名", ""))
        ws.cell(r, 15, f.get("meter_len") or f.get("采购米长", ""))
        ws.cell(r, 16, f.get("meter_price") or f.get("采购金额", ""))
        ws.cell(r, 17, designer)
        ws.cell(r, 18, f.get("arrived_time") or f.get("要求到货时间", ""))

    # === 辅料 (rows 32-43) ===
    trims = data.get("trims") or data.get("辅料信息", [])
    for idx, t in enumerate(trims):
        r = 32 + idx
        if r > 43:
            break
        ws.cell(r, 1, idx + 1)
        ws.cell(r, 2, t.get("name") or t.get("辅料名称", ""))
        ws.cell(r, 3, t.get("spec") or t.get("规格型号", ""))
        ws.cell(r, 4, t.get("color") or t.get("颜色", ""))
        ws.cell(r, 5, t.get("usage") or t.get("用量", ""))
        ws.cell(r, 6, t.get("unit") or t.get("单位", ""))
        ws.cell(r, 7, t.get("price") or t.get("采购单价", ""))
        ws.cell(r, 8, t.get("supplier") or t.get("供应商", ""))
        ws.cell(r, 9, t.get("phone") or t.get("联系方式", ""))
        ws.cell(r, 10, t.get("purchase_qty") or t.get("采购数量", ""))
        ws.cell(r, 11, designer)
        ws.cell(r, 12, t.get("arrived_time") or t.get("要求到货时间", ""))

    wb.save(output_path)
    return output_path

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python excel_writer.py <data.json> <output.xlsx>")
        sys.exit(1)
    with open(sys.argv[1], "r", encoding="utf-8") as f:
        data = json.load(f)
    path = build_excel(data, sys.argv[2])
    print(f"OK:{path}")
